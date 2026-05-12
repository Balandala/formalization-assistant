import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK


SERVICE_DIR = (
    Path(__file__).resolve().parents[1] / "app" / "services" / "formating_lib"
)
if str(SERVICE_DIR) not in sys.path:
    sys.path.insert(0, str(SERVICE_DIR))

from formating_config import Config  # noqa: E402
from word_processor import WordProcessor  # noqa: E402


def _create_document(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Титульный лист")
    doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph("1 Введение")
    doc.add_paragraph("Текст раздела.")
    doc.add_paragraph("2 Анализ")
    doc.save(path)


def test_process_file_check_only_keeps_original_file(tmp_path):
    filepath = tmp_path / "sample.docx"
    _create_document(filepath)
    before = filepath.read_bytes()

    report = WordProcessor(
        Config(table_of_contents=True, table_of_contents_page=2)
    ).process_file(str(filepath), check_only=True)

    assert filepath.read_bytes() == before
    assert report["check_only"] is True
    assert report["headings_detected"] == 2
    assert report["page_numbering_added"] is True
    assert report["page_fields_set"] is True
    assert report["table_of_contents_generated"] is True
    assert report["table_of_contents_entries"] == 2
    assert report["table_of_contents_page"] == 2
    assert report["diff_available"] is True
    assert report["diff_has_changes"] is True
    assert Path(report["diff_path"]).exists()


def test_process_file_without_check_only_updates_document(tmp_path):
    filepath = tmp_path / "sample.docx"
    _create_document(filepath)
    before = filepath.read_bytes()

    report = WordProcessor(
        Config(table_of_contents=True, table_of_contents_page=2)
    ).process_file(str(filepath))

    assert filepath.read_bytes() != before
    assert report["check_only"] is False
    assert report["headings_detected"] == 2
    assert report["table_of_contents_generated"] is True
    assert report["table_of_contents_entries"] == 2
    assert report["table_of_contents_page"] == 2
    assert report["diff_available"] is True
    assert report["diff_has_changes"] is True
    assert Path(report["diff_path"]).exists()

    with zipfile.ZipFile(filepath) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert "Содержание" in document_xml
    assert "1 Введение" in document_xml
    assert "2 Анализ" in document_xml
    assert "Обновите поле содержания" not in document_xml


def test_diff_file_contains_expected_labels(tmp_path):
    filepath = tmp_path / "sample.docx"
    _create_document(filepath)

    report = WordProcessor(
        Config(table_of_contents=True, table_of_contents_page=2)
    ).process_file(str(filepath), check_only=True)
    diff_html = Path(report["diff_path"]).read_text(encoding="utf-8")

    assert "Diff документа" in diff_html
    assert "Исходный документ" in diff_html
    assert "Потенциальный результат" in diff_html
    assert "Heading" in diff_html
    assert "Normal" in diff_html
    assert "Содержание" in diff_html


def test_process_file_without_toc_does_not_insert_contents(tmp_path):
    filepath = tmp_path / "sample.docx"
    _create_document(filepath)

    report = WordProcessor(Config(table_of_contents=False)).process_file(str(filepath))

    assert report["table_of_contents_generated"] is False
    assert report["table_of_contents_entries"] == 0


def test_process_file_adds_page_breaks_before_detected_headings(tmp_path):
    filepath = tmp_path / "heading-breaks.docx"
    doc = Document()
    doc.add_paragraph("Вступление.")
    doc.add_paragraph("1 Введение")
    doc.add_paragraph("Текст раздела.")
    doc.add_paragraph("2 Анализ")
    doc.save(filepath)

    report = WordProcessor(Config(table_of_contents=False)).process_file(str(filepath))

    assert report["headings_detected"] == 2
    assert report["heading_page_breaks_added"] == 2

    with zipfile.ZipFile(filepath) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert document_xml.count('w:type="page"') >= 2


def test_check_only_reports_heading_page_breaks(tmp_path):
    filepath = tmp_path / "heading-breaks.docx"
    doc = Document()
    doc.add_paragraph("Вступление.")
    doc.add_paragraph("1 Введение")
    doc.add_paragraph("Текст раздела.")
    doc.add_paragraph("2 Анализ")
    doc.save(filepath)
    before = filepath.read_bytes()

    report = WordProcessor(Config(table_of_contents=False)).process_file(
        str(filepath), check_only=True
    )

    assert filepath.read_bytes() == before
    assert report["heading_page_breaks_added"] == 2
    assert any(
        "Добавлен разрыв страницы перед заголовком" in detail
        for detail in report["details"]
    )
